"""File extractor tests covering each format we support without LibreOffice."""

import csv
import io
from pathlib import Path

import pytest


def _write(tmp_path, name, content_bytes):
    p = tmp_path / name
    p.write_bytes(content_bytes)
    return p


def test_extract_txt(tmp_path):
    from app.extractors import extract
    p = _write(tmp_path, "x.txt", "Jane Smith\njane@x.org".encode())
    r = extract(p)
    assert "Jane Smith" in r.text
    assert r.original_suffix == "txt"
    assert r.output_ext == ".txt"


def test_extract_csv(tmp_path):
    from app.extractors import extract
    out = io.StringIO()
    csv.writer(out).writerows([["name", "email"], ["Jane", "jane@x.org"]])
    p = _write(tmp_path, "x.csv", out.getvalue().encode())
    r = extract(p)
    assert "Jane" in r.text
    assert r.payload["rows"][0] == ["name", "email"]
    assert r.output_ext == ".csv"


def test_extract_html_strips_tags(tmp_path):
    from app.extractors import extract
    body = "<html><body><p>Hello <b>Jane</b></p><script>x()</script></body></html>"
    p = _write(tmp_path, "x.html", body.encode())
    r = extract(p)
    assert "Jane" in r.text
    assert "<b>" not in r.text
    assert "x()" not in r.text


def test_extract_xlsx_roundtrip(tmp_path):
    from openpyxl import Workbook
    from app.extractors import extract

    wb = Workbook()
    ws = wb.active
    ws.title = "people"
    ws.append(["name", "email"])
    ws.append(["Jane Smith", "jane@x.org"])
    p = tmp_path / "donors.xlsx"
    wb.save(str(p))

    r = extract(p)
    assert "Jane Smith" in r.text
    assert "jane@x.org" in r.text
    assert r.sheet_count == 1
    assert r.output_ext == ".xlsx"


def test_extract_docx_roundtrip(tmp_path):
    from docx import Document
    from app.extractors import extract

    doc = Document()
    doc.add_paragraph("From Jane Smith")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "email"
    table.rows[0].cells[1].text = "jane@x.org"
    p = tmp_path / "memo.docx"
    doc.save(str(p))

    r = extract(p)
    assert "Jane Smith" in r.text
    assert "jane@x.org" in r.text
    assert r.output_ext == ".docx"


def test_extract_pptx_roundtrip(tmp_path):
    from pptx import Presentation
    from app.extractors import extract

    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Jane Smith Q3"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "donor: jane@x.org"
    p = tmp_path / "deck.pptx"
    pres.save(str(p))

    r = extract(p)
    assert "Jane Smith" in r.text
    assert "jane@x.org" in r.text  # speaker notes captured
    assert r.output_ext == ".txt"


def test_extract_unsupported_raises(tmp_path):
    from app.extractors import ExtractError, extract
    p = _write(tmp_path, "x.foo", b"data")
    with pytest.raises(ExtractError):
        extract(p)
