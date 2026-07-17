"""Scrubber + deep scrub + verifier tests."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path


def _make_xlsx(path: Path, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    # Set author metadata so we can confirm deep scrub clears it.
    wb.properties.creator = "Jane Smith"
    wb.properties.lastModifiedBy = "Jane Smith"
    wb.save(str(path))


def _make_docx(path: Path, paragraphs, *, author="Jane Smith"):
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    doc.save(str(path))


def test_scrub_text_replaces(tmp_path):
    from app.extractors import extract
    from app.scrubber import scrub_text

    src = tmp_path / "x.txt"
    src.write_text("Jane Smith emailed jane@x.org")
    extracted = extract(src)
    rmap = {"Jane Smith": "[PERSON_3A4F]", "jane@x.org": "[EMAIL_3A4F]"}
    out = tmp_path / "x_anon.txt"
    result = scrub_text(extracted, rmap, out)
    text = out.read_text()
    assert "Jane Smith" not in text
    assert "jane@x.org" not in text
    assert "[PERSON_3A4F]" in text
    assert result.bytes_written > 0


def test_scrub_csv_replaces_cells(tmp_path):
    from app.extractors import extract
    from app.scrubber import scrub_csv

    src = tmp_path / "donors.csv"
    src.write_text("name,email\nJane Smith,jane@x.org\n")
    extracted = extract(src)
    rmap = {"Jane Smith": "[PERSON_3A4F]", "jane@x.org": "[EMAIL_3A4F]"}
    out = tmp_path / "donors_anon.csv"
    scrub_csv(extracted, rmap, out)
    body = out.read_text()
    assert "Jane Smith" not in body
    assert "[PERSON_3A4F]" in body


def test_scrub_xlsx_replaces_and_clears_metadata(tmp_path):
    from openpyxl import load_workbook
    from app.scrubber import scrub_xlsx

    src = tmp_path / "donors.xlsx"
    _make_xlsx(src, [["name", "email"], ["Jane Smith", "jane@x.org"]])

    rmap = {"Jane Smith": "[PERSON_3A4F]", "jane@x.org": "[EMAIL_3A4F]"}
    out = tmp_path / "donors_anon.xlsx"
    result = scrub_xlsx(src, rmap, out)

    # Cells replaced
    wb = load_workbook(str(out))
    ws = wb.active
    cells = [c.value for row in ws.iter_rows() for c in row]
    assert "Jane Smith" not in cells
    assert "[PERSON_3A4F]" in cells

    # Author cleared
    assert (wb.properties.creator or "") == ""
    assert (wb.properties.lastModifiedBy or "") == ""

    # Verify ZIP-level deep scrub stripped Jane Smith from raw XML.
    with zipfile.ZipFile(out) as zf:
        for name in zf.namelist():
            data = zf.read(name)
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                continue
            assert "Jane Smith" not in text, f"residual in {name}"
            assert "jane@x.org" not in text, f"residual in {name}"


def test_scrub_docx_replaces_and_strips_metadata(tmp_path):
    from docx import Document
    from app.scrubber import scrub_docx

    src = tmp_path / "memo.docx"
    _make_docx(src, ["From Jane Smith", "Email jane@x.org"])

    rmap = {"Jane Smith": "[PERSON_3A4F]", "jane@x.org": "[EMAIL_3A4F]"}
    out = tmp_path / "memo_anon.docx"
    scrub_docx(src, rmap, out)

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Jane Smith" not in text
    assert "[PERSON_3A4F]" in text
    assert (doc.core_properties.author or "") == ""
    assert (doc.core_properties.last_modified_by or "") == ""

    # Raw ZIP deep scrub
    with zipfile.ZipFile(out) as zf:
        for name in zf.namelist():
            data = zf.read(name)
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                continue
            assert "Jane Smith" not in text, f"residual in {name}"


def test_scrub_docx_strips_tracked_changes(tmp_path):
    """A DOCX with a manually injected <w:ins>/<w:del> region should be scrubbed clean."""
    import shutil
    from app.scrubber import scrub_docx

    src = tmp_path / "memo.docx"
    _make_docx(src, ["Hello world"])

    # Inject a fake tracked-change region holding "Jane Smith".
    with zipfile.ZipFile(src, "r") as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    doc_xml = parts["word/document.xml"].decode("utf-8")
    injection = (
        '<w:ins w:id="99" w:author="Jane Smith" w:date="2026-05-01T00:00:00Z">'
        '<w:r><w:t>Jane Smith</w:t></w:r></w:ins>'
        '<w:del w:id="100" w:author="Jane Smith" w:date="2026-05-01T00:00:00Z">'
        '<w:r><w:delText>Jane Smith</w:delText></w:r></w:del>'
    )
    doc_xml = doc_xml.replace("</w:body>", injection + "</w:body>")
    parts["word/document.xml"] = doc_xml.encode("utf-8")

    src2 = tmp_path / "memo_with_tracked.docx"
    with zipfile.ZipFile(src2, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)

    rmap = {"Jane Smith": "[PERSON_3A4F]"}
    out = tmp_path / "memo_anon.docx"
    scrub_docx(src2, rmap, out)

    with zipfile.ZipFile(out) as zf:
        body = zf.read("word/document.xml").decode("utf-8")
    assert "Jane Smith" not in body
    # All tracked-change wrappers stripped.
    assert "<w:ins" not in body
    assert "<w:del" not in body


def test_verifier_passes_clean_output(tmp_path):
    from app.scrubber import scrub_text
    from app.extractors import extract
    from app.verifier import verify_output

    src = tmp_path / "x.txt"
    src.write_text("Jane Smith and jane@x.org")
    rmap = {"Jane Smith": "[PERSON_3A4F]", "jane@x.org": "[EMAIL_3A4F]"}
    out = tmp_path / "anon.txt"
    scrub_text(extract(src), rmap, out)
    result = verify_output(out, rmap)
    assert result.passed is True
    assert result.total_matches == 0


def test_verifier_flags_residual_pii(tmp_path):
    """If we 'forget' a value in the map but it's still in the file, verify must fail."""
    from app.verifier import verify_output

    out = tmp_path / "leaky.txt"
    out.write_text("Jane Smith remains")
    # Map says we replaced Jane Smith but the file still contains it.
    rmap = {"Jane Smith": "[PERSON_3A4F]"}
    result = verify_output(out, rmap)
    assert result.passed is False
    assert "PERSON" in result.map_match_types


def test_verifier_regex_net_warns_without_blocking(tmp_path):
    """CODE_REVIEW C1: generic pattern hits are warnings, not hard failures.

    The regexes can't tell a missed email from any address-shaped string, so
    they surface for operator review instead of dead-ending the run.
    """
    from app.verifier import verify_output

    out = tmp_path / "leaky.txt"
    out.write_text("contact rogue@external.com")  # no map entry
    result = verify_output(out, {})
    assert result.passed is True
    assert "EMAIL" in result.regex_match_types  # surfaced as a warning


def test_verifier_no_false_positive_block_on_clean_doc(tmp_path):
    """CODE_REVIEW C1 repro: invoice numbers and version strings must not
    block a clean document, and must not even warn."""
    from app.verifier import verify_output

    out = tmp_path / "clean.txt"
    out.write_text("Invoice 5124403981. Version 1.2.3.4 released. 100 units.")
    result = verify_output(out, {"John Smith": "[PERSON_AAAA]"})
    assert result.passed is True
    assert result.regex_match_types == []


def test_verifier_deep_scans_xlsx_formula_literals(tmp_path):
    """CODE_REVIEW H3: PII living only inside a formula literal is invisible
    to the cell extractor (data_only) but must still fail the map scan."""
    from openpyxl import Workbook
    from app.verifier import verify_output

    leaky = tmp_path / "leaky.xlsx"
    wb = Workbook()
    wb.active["A1"] = '=CONCAT("Jane Smith", "!")'
    wb.save(str(leaky))

    result = verify_output(leaky, {"Jane Smith": "[PERSON_3A4F]"})
    assert result.passed is False
    assert "PERSON" in result.map_match_types


def test_scrub_docx_replaces_split_runs(tmp_path):
    """CODE_REVIEW C3 repro: Word splits strings across runs after edits.
    The run-aware pass must still replace them, formatting preserved."""
    from docx import Document
    from app.scrubber import scrub_docx

    src = tmp_path / "split.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Contact Ja")
    r = p.add_run("ne Smith today")
    r.bold = True
    doc.save(str(src))

    rmap = {"Jane Smith": "[PERSON_BBBB]"}
    out = tmp_path / "split_anon.docx"
    scrub_docx(src, rmap, out)

    result = Document(str(out))
    text = "\n".join(pp.text for pp in result.paragraphs)
    assert "Jane Smith" not in text
    assert "[PERSON_BBBB]" in text
    # Runs outside the match keep their formatting.
    assert any(run.bold for pp in result.paragraphs for run in pp.runs if "today" in run.text)


def test_verifier_ignores_placeholders(tmp_path):
    """Regex shouldn't trip on our own placeholder format."""
    from app.verifier import verify_output

    out = tmp_path / "clean.txt"
    out.write_text("contact [EMAIL_3A4F] for updates")
    result = verify_output(out, {})
    assert result.passed is True


def test_xlsx_formula_warning_collected(tmp_path):
    """A formula referencing a PII string should be flagged but not silently rewritten."""
    from openpyxl import Workbook, load_workbook
    from app.scrubber import scrub_xlsx

    src = tmp_path / "f.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Jane Smith"
    ws["A2"] = '=CONCAT("Jane Smith", "!")'
    wb.save(str(src))

    rmap = {"Jane Smith": "[PERSON_3A4F]"}
    out = tmp_path / "f_anon.xlsx"
    result = scrub_xlsx(src, rmap, out)

    assert any("formula references PII" in w for w in result.formula_warnings)

    # Non-formula cell IS scrubbed.
    wb2 = load_workbook(str(out))
    assert wb2.active["A1"].value == "[PERSON_3A4F]"
    # Formula cell preserved as-is.
    assert wb2.active["A2"].value.startswith("=")
